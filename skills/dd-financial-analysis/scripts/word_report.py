"""Render the controlled financial report JSON into a professional DOCX."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any


MARKDOWN_BOLD_RE = re.compile(r'(\*\*[^*]+\*\*|`[^`]+`)')
TABLE_SEPARATOR_RE = re.compile(r'^:?-{3,}:?$')
LATIN_FONT = 'Arial'
CJK_FONT = 'Hiragino Sans GB'


def _imports() -> dict[str, Any]:
    try:
        from docx import Document
        from docx.enum.section import WD_SECTION
        from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt, RGBColor
    except ImportError as exc:
        raise RuntimeError(
            '生成 Word 需要 python-docx；请使用 Codex workspace dependencies 返回的 Python 运行 render.py'
        ) from exc
    return locals()


def _set_run_font(run: Any, api: dict[str, Any], *, size: float | None = None,
                  bold: bool | None = None, color: str | None = None) -> None:
    qn = api['qn']
    Pt = api['Pt']
    RGBColor = api['RGBColor']
    run.font.name = LATIN_FONT
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn('w:ascii'), LATIN_FONT)
    fonts.set(qn('w:hAnsi'), LATIN_FONT)
    fonts.set(qn('w:cs'), LATIN_FONT)
    fonts.set(qn('w:eastAsia'), CJK_FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _set_style_font(style: Any, api: dict[str, Any]) -> None:
    qn = api['qn']
    style.font.name = LATIN_FONT
    fonts = style._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn('w:ascii'), LATIN_FONT)
    fonts.set(qn('w:hAnsi'), LATIN_FONT)
    fonts.set(qn('w:cs'), LATIN_FONT)
    fonts.set(qn('w:eastAsia'), CJK_FONT)


def _shade(element: Any, api: dict[str, Any], fill: str) -> None:
    OxmlElement = api['OxmlElement']
    qn = api['qn']
    properties = element.get_or_add_tcPr() if hasattr(element, 'get_or_add_tcPr') else element
    shd = properties.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        properties.append(shd)
    shd.set(qn('w:fill'), fill)


def _set_cell_margins(cell: Any, api: dict[str, Any], top: int = 80, start: int = 120,
                      bottom: int = 80, end: int = 120) -> None:
    OxmlElement = api['OxmlElement']
    qn = api['qn']
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for tag, value in (('top', top), ('start', start), ('bottom', bottom), ('end', end)):
        node = tc_mar.find(qn(f'w:{tag}'))
        if node is None:
            node = OxmlElement(f'w:{tag}')
            tc_mar.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')


def _set_table_geometry(table: Any, api: dict[str, Any], widths: list[int]) -> None:
    OxmlElement = api['OxmlElement']
    qn = api['qn']
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in('w:tblW')
    if tbl_w is None:
        tbl_w = OxmlElement('w:tblW')
        tbl_pr.append(tbl_w)
    tbl_w.set(qn('w:w'), str(sum(widths)))
    tbl_w.set(qn('w:type'), 'dxa')
    tbl_ind = tbl_pr.first_child_found_in('w:tblInd')
    if tbl_ind is None:
        tbl_ind = OxmlElement('w:tblInd')
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn('w:w'), '120')
    tbl_ind.set(qn('w:type'), 'dxa')

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement('w:gridCol')
        col.set(qn('w:w'), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths[min(index, len(widths) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in('w:tcW')
            if tc_w is None:
                tc_w = OxmlElement('w:tcW')
                tc_pr.append(tc_w)
            tc_w.set(qn('w:w'), str(width))
            tc_w.set(qn('w:type'), 'dxa')
            _set_cell_margins(cell, api)


def _column_widths(column_count: int, rows: list[list[str]]) -> list[int]:
    total = 9360
    if column_count <= 1:
        return [total]
    if column_count == 2:
        return [2100, 7260]
    if column_count == 3:
        return [1700, 3880, 3780]
    if column_count == 4:
        return [2300, 2353, 2353, 2354]
    if column_count >= 5:
        first = 2200 if any(len(row[0]) > 8 for row in rows if row) else 1700
        remainder = total - first
        base = remainder // (column_count - 1)
        widths = [first] + [base] * (column_count - 1)
        widths[-1] += total - sum(widths)
        return widths
    base = total // column_count
    widths = [base] * column_count
    widths[-1] += total - sum(widths)
    return widths


def _add_inline(paragraph: Any, text: str, api: dict[str, Any], *, size: float = 11,
                color: str = '202124') -> None:
    cursor = 0
    for match in MARKDOWN_BOLD_RE.finditer(text):
        if match.start() > cursor:
            _set_run_font(paragraph.add_run(text[cursor:match.start()]), api, size=size, color=color)
        token = match.group(0)
        if token.startswith('**'):
            run = paragraph.add_run(token[2:-2])
            _set_run_font(run, api, size=size, bold=True, color=color)
        else:
            run = paragraph.add_run(token[1:-1])
            _set_run_font(run, api, size=max(size - 0.5, 8), color='44546A')
        cursor = match.end()
    if cursor < len(text):
        _set_run_font(paragraph.add_run(text[cursor:]), api, size=size, color=color)


def _is_separator_row(cells: list[str]) -> bool:
    return bool(cells) and all(TABLE_SEPARATOR_RE.fullmatch(cell.replace(' ', '')) for cell in cells)


def _split_table_line(line: str) -> list[str]:
    text = line.strip().strip('|')
    cells: list[str] = []
    current: list[str] = []
    escaped = False
    for char in text:
        if escaped:
            current.append(char)
            escaped = False
        elif char == '\\':
            escaped = True
        elif char == '|':
            cells.append(''.join(current).strip())
            current = []
        else:
            current.append(char)
    if escaped:
        current.append('\\')
    cells.append(''.join(current).strip())
    return cells


def _parse_table(lines: list[str]) -> list[list[str]]:
    rows = []
    for line in lines:
        cells = _split_table_line(line)
        if not _is_separator_row(cells):
            rows.append(cells)
    return rows


def _add_table(doc: Any, rows: list[list[str]], api: dict[str, Any]) -> None:
    if not rows:
        return
    WD_TABLE_ALIGNMENT = api['WD_TABLE_ALIGNMENT']
    WD_CELL_VERTICAL_ALIGNMENT = api['WD_CELL_VERTICAL_ALIGNMENT']
    WD_ALIGN_PARAGRAPH = api['WD_ALIGN_PARAGRAPH']
    OxmlElement = api['OxmlElement']
    qn = api['qn']
    column_count = max(len(row) for row in rows)
    normalized = [row + [''] * (column_count - len(row)) for row in rows]
    table = doc.add_table(rows=len(normalized), cols=column_count)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = 'Table Grid'
    widths = _column_widths(column_count, normalized)
    _set_table_geometry(table, api, widths)
    for row_index, values in enumerate(normalized):
        row = table.rows[row_index]
        cant_split = OxmlElement('w:cantSplit')
        row._tr.get_or_add_trPr().append(cant_split)
        if row_index == 0:
            header = OxmlElement('w:tblHeader')
            header.set(qn('w:val'), 'true')
            row._tr.get_or_add_trPr().append(header)
        for col_index, value in enumerate(values):
            cell = row.cells[col_index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index == 0:
                _shade(cell._tc, api, 'E8EEF5')
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_before = api['Pt'](0)
            paragraph.paragraph_format.space_after = api['Pt'](0)
            paragraph.paragraph_format.line_spacing = 1.0
            compact = len(value) <= 24
            if col_index > 0 and compact:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            _add_inline(paragraph, value, api, size=8.5, color='202124')
            if row_index == 0:
                for run in paragraph.runs:
                    run.bold = True
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = api['Pt'](2)


def _add_callout(doc: Any, text: str, api: dict[str, Any], *, caution: bool = False) -> None:
    OxmlElement = api['OxmlElement']
    qn = api['qn']
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = api['Pt'](5)
    paragraph.paragraph_format.space_after = api['Pt'](8)
    paragraph.paragraph_format.left_indent = api['Inches'](0.12)
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), 'FFF4CE' if caution else 'F2F4F7')
    p_pr.append(shd)
    borders = OxmlElement('w:pBdr')
    left = OxmlElement('w:left')
    left.set(qn('w:val'), 'single')
    left.set(qn('w:sz'), '18')
    left.set(qn('w:space'), '6')
    left.set(qn('w:color'), 'C99A00' if caution else '4472C4')
    borders.append(left)
    p_pr.append(borders)
    _add_inline(paragraph, text, api, size=10.5, color='3A2E00' if caution else '203864')


def _missing_material_categories(items: list[dict[str, Any]]) -> str:
    categories: list[str] = []
    for item in items:
        category = str(item.get('category') or '').strip()
        if category and category not in categories:
            categories.append(category)
    return '、'.join(categories)


def _source_origin_label(value: Any) -> str:
    return {
        'provided': '进件材料',
        'existing-public-file': '已提供公开资料',
        'public-research': '联网补充',
    }.get(str(value or ''), str(value or '未标明'))


def _review_item_text(item: dict[str, Any]) -> str:
    details = [str(value) for value in (item.get('details') or []) if str(value).strip()]
    impact = str(item.get('impact') or '').strip()
    if details and impact:
        return f'{"、".join(details)}：{impact}'
    return '、'.join(details) or impact or '未明确'


def _add_review_register_section(
    doc: Any,
    report: dict[str, Any],
    api: dict[str, Any],
) -> None:
    missing_materials = [
        item for item in (report.get('missingMaterials') or []) if isinstance(item, dict)
    ]
    register = report.get('reviewRegister') or {}
    register_items = [
        item for item in (register.get('items') or []) if isinstance(item, dict)
    ]
    sources = [item for item in (register.get('sources') or []) if isinstance(item, dict)]
    provenance = report.get('_provenance') or {}
    heading = doc.add_paragraph(style='Heading 1')
    heading.paragraph_format.keep_with_next = True
    _add_inline(heading, '资料与核验事项', api, size=16, color='2E5B88')
    lead = doc.add_paragraph()
    lead.paragraph_format.space_after = api['Pt'](8)
    lead.paragraph_format.line_spacing = 1.10
    _add_inline(
        lead,
        '正文已按现有材料完成。以下事项集中说明结论边界、待补资料和采用来源。',
        api,
        size=10.5,
    )

    if missing_materials:
        subheading = doc.add_paragraph(style='Heading 2')
        _add_inline(subheading, '待补资料', api, size=13, color='2E5B88')
        rows = [['类别', '建议补充资料', '当前影响']]
        for item in missing_materials:
            rows.append([
                str(item.get('category') or '未分类'),
                str(item.get('required') or '未明确'),
                str(item.get('impact') or '未明确'),
            ])
        _add_table(doc, rows, api)

    other_items = [
        item for item in register_items
        if item.get('kind') != 'missing-evidence' and item.get('audience') == 'report'
    ]
    if other_items:
        subheading = doc.add_paragraph(style='Heading 2')
        _add_inline(subheading, '其他核验事项', api, size=13, color='2E5B88')
        rows = [['类别', '事项及影响', '处理']]
        for item in other_items:
            rows.append([
                str(item.get('category') or '未分类'),
                _review_item_text(item),
                str(item.get('requiredAction') or '未明确'),
            ])
        _add_table(doc, rows, api)

    if sources:
        subheading = doc.add_paragraph(style='Heading 2')
        _add_inline(subheading, '采用资料', api, size=13, color='2E5B88')
        rows = [['资料', '期间', '来源', '定位']]
        for source in sources:
            publisher = str(source.get('publisher') or '').strip()
            origin = _source_origin_label(source.get('origin'))
            rows.append([
                str(source.get('displayName') or '未命名资料'),
                str(source.get('period') or '—'),
                f'{origin}{" / " + publisher if publisher else ""}',
                str(source.get('locator') or '—'),
            ])
        _add_table(doc, rows, api)

    subheading = doc.add_paragraph(style='Heading 2')
    _add_inline(subheading, '生成追溯', api, size=13, color='2E5B88')
    trace = [
        ['报告状态', '有限材料' if provenance.get('reportStatus') == 'limited' else '已核验'],
        ['证据覆盖', {
            'partial': '部分', 'complete': '完整',
        }.get(provenance.get('evidenceCoverage'), '未记录')],
        ['渲染器版本', str(provenance.get('version') or '未记录')],
        ['最终计算摘要', str(provenance.get('calculationDigest') or '未记录')],
        ['核验事项摘要', str(provenance.get('reviewRegisterDigest') or '未记录')],
        ['公开研究摘要', str(provenance.get('publicResearchDigest') or '未记录')],
        ['运行标识', str(provenance.get('runId') or '未记录')],
    ]
    _add_table(doc, trace, api)
    _add_callout(
        doc,
        '详细检索记录、字段映射和完整计算结果见同名追溯 JSON。',
        api,
        caution=provenance.get('reportStatus') == 'limited',
    )


def _add_markdown(doc: Any, markdown: str, api: dict[str, Any]) -> None:
    lines = markdown.splitlines()
    index = 0
    while index < len(lines):
        raw = lines[index]
        stripped = raw.strip()
        if not stripped:
            index += 1
            continue
        if stripped.startswith('|'):
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith('|'):
                table_lines.append(lines[index].strip())
                index += 1
            _add_table(doc, _parse_table(table_lines), api)
            continue
        if stripped.startswith('#### '):
            paragraph = doc.add_paragraph(style='Heading 3')
            _add_inline(paragraph, stripped[5:], api, size=12, color='1F4D78')
        elif stripped.startswith('### '):
            paragraph = doc.add_paragraph(style='Heading 3')
            _add_inline(paragraph, stripped[4:], api, size=12, color='1F4D78')
        elif stripped.startswith('> '):
            _add_callout(doc, stripped[2:], api, caution=True)
        elif stripped.startswith('- '):
            paragraph = doc.add_paragraph(style='List Bullet')
            _add_inline(paragraph, stripped[2:], api, size=10.5)
        else:
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.space_after = api['Pt'](6)
            paragraph.paragraph_format.line_spacing = 1.10
            _add_inline(paragraph, stripped, api, size=10.5)
        index += 1


def _page_number(paragraph: Any, api: dict[str, Any]) -> None:
    OxmlElement = api['OxmlElement']
    qn = api['qn']
    run = paragraph.add_run()
    begin = OxmlElement('w:fldChar')
    begin.set(qn('w:fldCharType'), 'begin')
    instruction = OxmlElement('w:instrText')
    instruction.set(qn('xml:space'), 'preserve')
    instruction.text = ' PAGE '
    end = OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'), 'end')
    run._r.extend([begin, instruction, end])
    _set_run_font(run, api, size=9, color='6B7280')


def _configure_document(doc: Any, report: dict[str, Any], api: dict[str, Any]) -> None:
    WD_ALIGN_PARAGRAPH = api['WD_ALIGN_PARAGRAPH']
    Inches = api['Inches']
    Pt = api['Pt']
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.82)
    section.bottom_margin = Inches(0.78)
    section.left_margin = Inches(0.88)
    section.right_margin = Inches(0.88)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles['Normal']
    _set_style_font(normal, api)
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    for name, size, color, before, after in (
        ('Title', 24, '203864', 0, 8),
        ('Heading 1', 16, '2E5B88', 14, 8),
        ('Heading 2', 13, '2E5B88', 11, 6),
        ('Heading 3', 12, '1F4D78', 8, 4),
    ):
        style = styles[name]
        _set_style_font(style, api)
        style.font.size = Pt(size)
        style.font.color.rgb = api['RGBColor'].from_string(color)
        style.font.bold = name != 'Title'
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    for name in ('List Bullet', 'List Number'):
        style = styles[name]
        _set_style_font(style, api)
        style.font.size = Pt(10.5)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(4)

    header = section.header
    header_paragraph = header.paragraphs[0]
    header_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _set_run_font(
        header_paragraph.add_run(report.get('reportName') or '财务分析报告'),
        api, size=8.5, color='6B7280',
    )
    footer = section.footer
    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_run_font(footer_paragraph.add_run('第 '), api, size=9, color='6B7280')
    _page_number(footer_paragraph, api)
    _set_run_font(footer_paragraph.add_run(' 页'), api, size=9, color='6B7280')

    doc.core_properties.title = report.get('reportName') or '财务分析报告'
    doc.core_properties.subject = '对公信贷财务分析'
    doc.core_properties.author = 'dd-financial-analysis'


def write_word_report(
    report: dict[str, Any],
    result: dict[str, Any],
    output_path: str | Path,
) -> str:
    api = _imports()
    Document = api['Document']
    WD_ALIGN_PARAGRAPH = api['WD_ALIGN_PARAGRAPH']
    doc = Document()
    _configure_document(doc, report, api)
    provenance = report.get('_provenance') or {}
    verified = provenance.get('verified') is True
    limited_report = (
        provenance.get('reportStatus') == 'limited'
        or provenance.get('evidenceCoverage') == 'partial'
        or not verified
    )
    status = '有限材料财务分析报告' if limited_report else '已核验标准报告'
    missing_materials = [
        item for item in (report.get('missingMaterials') or []) if isinstance(item, dict)
    ]

    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_before = api['Pt'](54)
    kicker.paragraph_format.space_after = api['Pt'](12)
    _set_run_font(kicker.add_run('对公信贷'), api, size=11, bold=True, color='6B7280')

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = api['Pt'](10)
    _set_run_font(
        title.add_run(report.get('reportName') or '财务分析报告'),
        api, size=24, bold=True, color='203864',
    )
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = api['Pt'](22)
    _set_run_font(subtitle.add_run(status), api, size=13, bold=True,
                  color='8A5A00' if limited_report else '2E5B88')
    if missing_materials:
        categories = _missing_material_categories(missing_materials)
        _add_callout(
            doc,
            f'本报告已按现有材料完成分析。进一步核验仍需补充：{categories}；'
            '详见文末“资料与核验事项”。',
            api,
            caution=True,
        )
    else:
        _add_callout(
            doc,
            '本报告已按现有材料完成分析，财务数值由确定性脚本生成。'
            + ('具体核验边界见文末。' if limited_report else ''),
            api,
            caution=limited_report,
        )
    metadata = doc.add_paragraph()
    metadata.alignment = WD_ALIGN_PARAGRAPH.CENTER
    metadata.paragraph_format.space_before = api['Pt'](18)
    period = result.get('period_label') or '未标明期间'
    scope = (result.get('normalized_inputs') or {}).get('statement_scope') or '未明确口径'
    _set_run_font(metadata.add_run(f'报告期：{period}    报表口径：{scope}'), api,
                  size=10.5, color='4B5563')
    generated = doc.add_paragraph()
    generated.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(generated.add_run(f'生成时间：{report.get("generatedAt") or "未记录"}'), api,
                  size=9, color='6B7280')
    doc.add_page_break()

    for chapter in report.get('chapters') or []:
        heading = doc.add_paragraph(style='Heading 1')
        _add_inline(heading, str(chapter.get('title') or ''), api, size=16, color='2E5B88')
        _add_markdown(doc, str(chapter.get('content') or ''), api)
        for child in chapter.get('children') or []:
            child_heading = doc.add_paragraph(style='Heading 2')
            _add_inline(child_heading, str(child.get('title') or ''), api, size=13, color='2E5B88')
            _add_markdown(doc, str(child.get('content') or ''), api)

    _add_review_register_section(doc, report, api)

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    if not output.is_file() or output.stat().st_size == 0:
        raise RuntimeError('Word 报告写入失败')
    return str(output)
